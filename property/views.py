"""Property & Assets views. Every action gated server-side."""
import csv
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from .forms import (
    AssetForm, ConditionReportForm, NewItemApprovalForm, NewItemReportForm,
    ReceiptConditionForm, RecipientDecisionForm, RelocationApprovalForm,
    RelocationExceptionResolveForm, RelocationRequestForm, ReportFilterForm,
)
from .models import (
    Asset, AssetNewItemReport, AssetRelocation, ConditionReport,
    PhysicalArea, PropertyCheck, StaffAreaAccess,
)
from .permissions import (
    areas_for, is_active_staff, is_director, require_active_staff,
    require_director, user_can_view_asset,
)
from . import services


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------
@login_required
@require_active_staff
def dashboard(request):
    user = request.user
    user_areas = areas_for(user)
    is_dir = is_director(user)

    if is_dir:
        asset_qs = Asset.objects.all()
        check_qs = PropertyCheck.objects.all()
        cond_qs = ConditionReport.objects.all()
        new_item_qs = AssetNewItemReport.objects.filter(
            status=AssetNewItemReport.Status.PENDING)
        reloc_pending_dir = AssetRelocation.objects.filter(
            status=AssetRelocation.Status.PENDING_APPROVAL)
        reloc_awaiting_me = AssetRelocation.objects.filter(
            status__in=[AssetRelocation.Status.APPROVED,
                        AssetRelocation.Status.AWAITING_RECIPIENT])
        reloc_exceptions = AssetRelocation.objects.filter(
            status=AssetRelocation.Status.TRANSFER_EXCEPTION)
    else:
        asset_qs = Asset.objects.filter(current_area__code__in=user_areas)
        check_qs = PropertyCheck.objects.filter(area__code__in=user_areas)
        cond_qs = ConditionReport.objects.filter(area__code__in=user_areas)
        new_item_qs = AssetNewItemReport.objects.filter(
            reported_by=user, status=AssetNewItemReport.Status.PENDING)
        reloc_pending_dir = AssetRelocation.objects.none()
        reloc_awaiting_me = AssetRelocation.objects.filter(
            receiving_attendant=user,
            status__in=[AssetRelocation.Status.APPROVED,
                        AssetRelocation.Status.AWAITING_RECIPIENT])
        reloc_exceptions = AssetRelocation.objects.none()

    return render(request, "property/dashboard.html", {
        "is_dir": is_dir,
        "user_areas": sorted(user_areas),
        "asset_count": asset_qs.count(),
        "check_count": check_qs.count(),
        "open_conditions": cond_qs.exclude(
            status=ConditionReport.Status.RESOLVED).count(),
        "pending_new_items": new_item_qs.count(),
        "reloc_pending_dir": reloc_pending_dir.count(),
        "reloc_awaiting_me": reloc_awaiting_me.count(),
        "reloc_exceptions": reloc_exceptions.count(),
        "recent_assets": asset_qs.select_related("current_area").order_by("-updated_at")[:8],
        "recent_checks": check_qs.select_related("area", "checked_by").order_by("-checked_at")[:8],
        "recent_conditions": cond_qs.select_related("asset", "area", "reported_by").order_by("-reported_at")[:8],
    })


# ---------------------------------------------------------------------------
# Asset register
# ---------------------------------------------------------------------------
# Canonical display order for areas when grouping assets. This matches the
# order of the areas as they appear on the venue floor: Bar first, then
# kitchen, then dining / outdoor / facilities.
AREA_DISPLAY_ORDER = [
    "BAR", "BAR_COUNTER", "BUTCHERY", "CHOMA", "KITCHEN",
    "DINING_AREA", "OUTDOOR_AREA", "WASHROOMS", "GARDEN_COMPOUND",
    "OFFICE", "STORE", "OTHER",
]


@login_required
@require_active_staff
def asset_list(request):
    user = request.user
    qs = Asset.objects.select_related("current_area")
    if not is_director(user):
        qs = qs.filter(current_area__code__in=areas_for(user))
    area_code = request.GET.get("area") or ""
    if area_code:
        qs = qs.filter(current_area__code=area_code)
    q = request.GET.get("q") or ""
    if q:
        qs = qs.filter(Q(asset_number__icontains=q) | Q(name__icontains=q) |
                       Q(category__icontains=q) | Q(description__icontains=q))
    qs = qs.order_by("asset_number")

    # Group by area for the Director. Staff see a flat list (their scope is
    # typically one area, so grouping adds no value).
    grouped = None
    if is_director(user) and not q:
        by_code = {}
        for a in qs:
            by_code.setdefault(a.current_area.code, []).append(a)
        grouped = []
        for code in AREA_DISPLAY_ORDER:
            if code in by_code:
                # Look up the area object for the display name
                area_obj = PhysicalArea.objects.filter(code=code).first()
                if area_obj is None:
                    continue
                grouped.append({
                    "code": code,
                    "name": area_obj.name or code,
                    "assets": by_code[code],
                    "count": len(by_code[code]),
                })

    return render(request, "property/asset_list.html", {
        "assets": qs,
        "grouped": grouped,
        "areas": PhysicalArea.objects.filter(is_active=True).order_by("code"),
        "current_area": area_code,
        "q": q,
        "can_edit": is_director(user),
    })


@login_required
@require_active_staff
def asset_detail(request, pk):
    asset = get_object_or_404(Asset.objects.select_related("current_area"), pk=pk)
    if not user_can_view_asset(request.user, asset):
        raise PermissionDenied("No access to this asset.")
    return render(request, "property/asset_detail.html", {
        "asset": asset,
        "checks": asset.check_lines.select_related("property_check", "asset").order_by("-property_check__checked_at")[:20],
        "conditions": asset.condition_reports.select_related("reported_by", "resolved_by").order_by("-reported_at"),
        "relocations": asset.relocations.select_related("origin_area", "destination_area", "requested_by", "receiving_attendant").order_by("-submitted_at"),
        "can_edit": is_director(request.user),
    })


@login_required
@require_director
def asset_create(request):
    if request.method == "POST":
        form = AssetForm(request.POST)
        if form.is_valid():
            asset = services.create_asset(user=request.user, data=form.cleaned_data)
            messages.success(request, f"Asset {asset.asset_number} created.")
            return redirect("property:asset_detail", pk=asset.pk)
    else:
        form = AssetForm()
    return render(request, "property/asset_form.html", {"form": form, "mode": "create"})


@login_required
@require_director
def asset_edit(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    if request.method == "POST":
        form = AssetForm(request.POST, instance=asset)
        if form.is_valid():
            services.update_asset(user=request.user, asset=asset, data=form.cleaned_data)
            messages.success(request, f"Asset {asset.asset_number} updated.")
            return redirect("property:asset_detail", pk=asset.pk)
    else:
        form = AssetForm(instance=asset)
    return render(request, "property/asset_form.html", {"form": form, "asset": asset, "mode": "edit"})


@login_required
@require_director
def asset_dispose(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    if request.method != "POST":
        raise PermissionDenied("POST required.")
    reason = request.POST.get("reason", "").strip()
    try:
        services.dispose_asset(user=request.user, asset=asset, reason=reason)
        messages.success(request, f"Asset {asset.asset_number} disposed.")
    except ValidationError as e:
        messages.error(request, str(e))
    return redirect("property:asset_detail", pk=asset.pk)


# ---------------------------------------------------------------------------
# Property checks
# ---------------------------------------------------------------------------
@login_required
@require_active_staff
def check_list(request):
    qs = PropertyCheck.objects.select_related("area", "checked_by")
    if not is_director(request.user):
        qs = qs.filter(area__code__in=areas_for(request.user))
    return render(request, "property/check_list.html", {"checks": qs.order_by("-checked_at")[:200]})


@login_required
@require_active_staff
def check_create(request):
    user = request.user
    user_areas = areas_for(user)
    if request.method == "POST":
        area_code = request.POST.get("area", "")
        if area_code not in user_areas:
            raise PermissionDenied("Area not authorized.")
        area = get_object_or_404(PhysicalArea, code=area_code)
        lines = []
        idx = 0
        while f"asset_id_{idx}" in request.POST:
            asset_id = request.POST.get(f"asset_id_{idx}")
            if asset_id:
                lines.append({
                    "asset_id": int(asset_id),
                    "expected_quantity": request.POST.get(f"expected_{idx}", "0"),
                    "observed_quantity": request.POST.get(f"observed_{idx}", "0"),
                    "condition": request.POST.get(f"condition_{idx}", "WORKING"),
                    "notes": request.POST.get(f"notes_{idx}", ""),
                })
            idx += 1
        if not lines:
            messages.error(request, "Add at least one line.")
        else:
            try:
                check = services.submit_property_check(
                    user=user, area=area, lines=lines,
                    notes=request.POST.get("notes", ""),
                )
                messages.success(request, "Property check submitted.")
                return redirect("property:check_list")
            except (PermissionDenied, ValidationError) as e:
                messages.error(request, str(e))

    selected_area = request.GET.get("area") or (sorted(user_areas)[0] if user_areas else "")
    area_obj = PhysicalArea.objects.filter(code=selected_area).first() if selected_area else None
    assets = []
    if area_obj:
        assets = (Asset.objects.filter(current_area=area_obj,
                                       lifecycle_status=Asset.Lifecycle.ACTIVE)
                  .order_by("asset_number"))
    return render(request, "property/check_form.html", {
        "user_areas": sorted(user_areas),
        "selected_area": selected_area,
        "area_obj": area_obj,
        "assets": assets,
    })


# ---------------------------------------------------------------------------
# New-item reports
# ---------------------------------------------------------------------------
@login_required
@require_active_staff
def new_item_list(request):
    user = request.user
    if is_director(user):
        qs = AssetNewItemReport.objects.all()
    else:
        qs = AssetNewItemReport.objects.filter(reported_by=user)
    status_filter = request.GET.get("status") or ""
    if status_filter:
        qs = qs.filter(status=status_filter)
    return render(request, "property/new_item_report_list.html", {
        "reports": qs.select_related("area", "reported_by", "reviewed_by").order_by("-submitted_at"),
        "is_dir": is_director(user),
        "current_status": status_filter,
    })


@login_required
@require_active_staff
def new_item_create(request):
    user = request.user
    allowed = areas_for(user)
    if request.method == "POST":
        form = NewItemReportForm(request.POST, allowed_areas=allowed)
        if form.is_valid():
            report = services.submit_new_item_report(user=user, data=form.cleaned_data)
            messages.success(request, "Report submitted for Director review.")
            return redirect("property:new_item_list")
    else:
        form = NewItemReportForm(allowed_areas=allowed)
    return render(request, "property/new_item_report_form.html", {"form": form})


@login_required
@require_director
def new_item_review(request, pk):
    report = get_object_or_404(AssetNewItemReport, pk=pk)
    if request.method == "POST":
        form = NewItemApprovalForm(request.POST)
        if form.is_valid():
            try:
                if form.cleaned_data["action"] == "approve":
                    asset = services.approve_new_item_report(
                        user=request.user, report=report,
                        review_notes=form.cleaned_data["review_notes"],
                    )
                    messages.success(request, f"Approved. Created asset {asset.asset_number}.")
                else:
                    services.reject_new_item_report(
                        user=request.user, report=report,
                        review_notes=form.cleaned_data["review_notes"],
                    )
                    messages.success(request, "Report rejected.")
                return redirect("property:new_item_list")
            except (PermissionDenied, ValidationError) as e:
                messages.error(request, str(e))
    else:
        form = NewItemApprovalForm()
    return render(request, "property/new_item_review.html", {"report": report, "form": form})


# ---------------------------------------------------------------------------
# Condition reports
# ---------------------------------------------------------------------------
@login_required
@require_active_staff
def condition_list(request):
    qs = ConditionReport.objects.select_related("asset", "area", "reported_by", "resolved_by")
    if not is_director(request.user):
        qs = qs.filter(area__code__in=areas_for(request.user))
    status_filter = request.GET.get("status") or ""
    if status_filter == "open":
        qs = qs.exclude(status=ConditionReport.Status.RESOLVED)
    elif status_filter:
        qs = qs.filter(status=status_filter)
    return render(request, "property/condition_list.html", {
        "reports": qs.order_by("-reported_at")[:200],
        "is_dir": is_director(request.user),
        "current_status": status_filter,
    })


@login_required
@require_active_staff
def condition_create(request, asset_pk):
    asset = get_object_or_404(Asset, pk=asset_pk)
    if not user_can_view_asset(request.user, asset):
        raise PermissionDenied("No access to this asset.")
    if request.method == "POST":
        form = ConditionReportForm(request.POST)
        if form.is_valid():
            try:
                services.submit_condition_report(
                    user=request.user, asset=asset, data=form.cleaned_data,
                )
                messages.success(request, "Condition report submitted.")
                return redirect("property:condition_list")
            except (PermissionDenied, ValidationError) as e:
                messages.error(request, str(e))
    else:
        form = ConditionReportForm()
    return render(request, "property/condition_form.html", {"asset": asset, "form": form})


@login_required
@require_director
def condition_resolve(request, pk):
    report = get_object_or_404(ConditionReport, pk=pk)
    if request.method != "POST":
        raise PermissionDenied("POST required.")
    notes = request.POST.get("resolution_notes", "").strip()
    try:
        services.resolve_condition_report(user=request.user, report=report,
                                          resolution_notes=notes)
        messages.success(request, "Condition report resolved.")
    except (PermissionDenied, ValidationError) as e:
        messages.error(request, str(e))
    return redirect("property:condition_list")


# ---------------------------------------------------------------------------
# Relocation workflow
# ---------------------------------------------------------------------------
@login_required
@require_active_staff
def relocation_list(request):
    user = request.user
    if is_director(user):
        qs = AssetRelocation.objects.all()
    else:
        qs = AssetRelocation.objects.filter(
            Q(requested_by=user) | Q(receiving_attendant=user) |
            Q(origin_area__code__in=areas_for(user)) |
            Q(destination_area__code__in=areas_for(user))
        ).distinct()
    return render(request, "property/relocation_list.html", {
        "relocations": qs.select_related(
            "asset", "origin_area", "destination_area",
            "requested_by", "receiving_attendant", "approved_by"
        ).order_by("-submitted_at"),
        "is_dir": is_director(user),
    })


@login_required
@require_active_staff
def relocation_request(request, asset_pk):
    asset = get_object_or_404(Asset, pk=asset_pk)
    if not user_can_view_asset(request.user, asset):
        raise PermissionDenied("No access to this asset.")

    from .permissions import eligible_recipients_for_area
    all_areas = PhysicalArea.objects.filter(is_active=True).order_by("code")
    # Precompute {area_code: set(user_ids)} for the template
    eligible_map = {
        a.code: set(eligible_recipients_for_area(a.code).values_list("pk", flat=True))
        for a in all_areas
    }

    if request.method == "POST":
        form = RelocationRequestForm(request.POST, user=request.user, asset=asset)
        if form.is_valid():
            try:
                reloc = services.create_relocation_request(
                    user=request.user, asset=asset, data=form.cleaned_data,
                )
                messages.success(request, "Relocation request submitted for Director approval.")
                return redirect("property:relocation_detail", pk=reloc.pk)
            except (PermissionDenied, ValidationError) as e:
                messages.error(request, str(e))
    else:
        form = RelocationRequestForm(user=request.user, asset=asset)
    import json
    eligible_map_json = json.dumps(
        {code: sorted(list(ids)) for code, ids in eligible_map.items()}
    )
    destination_area_codes_json = json.dumps(
        {str(a.pk): a.code for a in all_areas}
    )
    return render(request, "property/relocation_request_form.html", {
        "asset": asset,
        "form": form,
        "eligible_map": eligible_map,
        "eligible_map_json": eligible_map_json,
        "destination_area_codes_json": destination_area_codes_json,
    })


@login_required
@require_active_staff
def relocation_detail(request, pk):
    reloc = get_object_or_404(AssetRelocation.objects.select_related(
        "asset", "origin_area", "destination_area",
        "requested_by", "receiving_attendant", "approved_by", "received_by",
    ), pk=pk)
    user = request.user
    can_see = (
        is_director(user) or reloc.requested_by_id == user.pk or
        reloc.receiving_attendant_id == user.pk or
        reloc.origin_area.code in areas_for(user) or
        reloc.destination_area.code in areas_for(user)
    )
    if not can_see:
        raise PermissionDenied("No access to this relocation.")
    return render(request, "property/relocation_detail.html", {
        "reloc": reloc,
        "is_dir": is_director(user),
        "is_recipient": reloc.receiving_attendant_id == user.pk,
        "approval_form": RelocationApprovalForm() if is_director(user) and reloc.is_pending_director else None,
        "recipient_form": RecipientDecisionForm() if reloc.is_awaiting_recipient and reloc.receiving_attendant_id == user.pk else None,
        "receipt_form": ReceiptConditionForm() if reloc.status == AssetRelocation.Status.RECEIVED or (reloc.recipient_decision == "ACCEPTED" and reloc.status in {AssetRelocation.Status.APPROVED, AssetRelocation.Status.AWAITING_RECIPIENT}) and reloc.receiving_attendant_id == user.pk else None,
    })


@login_required
@require_director
def approval_queue(request):
    pending = (AssetRelocation.objects
               .filter(status=AssetRelocation.Status.PENDING_APPROVAL)
               .select_related("asset", "origin_area", "destination_area",
                               "requested_by", "receiving_attendant")
               .order_by("submitted_at"))
    exceptions = (AssetRelocation.objects
                  .filter(status=AssetRelocation.Status.TRANSFER_EXCEPTION)
                  .select_related("asset", "origin_area", "destination_area",
                                  "requested_by", "receiving_attendant")
                  .order_by("-submitted_at"))
    return render(request, "property/approval_queue.html", {
        "pending": pending,
        "exceptions": exceptions,
    })


@login_required
@require_director
def relocation_decide(request, pk):
    reloc = get_object_or_404(AssetRelocation, pk=pk)
    if request.method != "POST":
        raise PermissionDenied("POST required.")
    form = RelocationApprovalForm(request.POST)
    if form.is_valid():
        try:
            services.director_decide_relocation(
                user=request.user, relocation=reloc,
                action=form.cleaned_data["action"],
                notes=form.cleaned_data["director_decision_notes"],
            )
            messages.success(request, f"Relocation {form.cleaned_data['action']}d.")
        except (PermissionDenied, ValidationError) as e:
            messages.error(request, str(e))
    return redirect("property:approval_queue")


@login_required
@require_active_staff
def relocation_recipient_decide(request, pk):
    reloc = get_object_or_404(AssetRelocation, pk=pk)
    if request.method != "POST":
        raise PermissionDenied("POST required.")
    form = RecipientDecisionForm(request.POST)
    if form.is_valid():
        try:
            services.recipient_decide(
                user=request.user, relocation=reloc,
                action=form.cleaned_data["action"],
                decline_reason=form.cleaned_data.get("decline_reason", ""),
            )
            messages.success(request, "Your decision was recorded.")
        except (PermissionDenied, ValidationError) as e:
            messages.error(request, str(e))
    return redirect("property:relocation_detail", pk=reloc.pk)


@login_required
@require_active_staff
def relocation_confirm_receipt(request, pk):
    reloc = get_object_or_404(AssetRelocation, pk=pk)
    if request.method != "POST":
        raise PermissionDenied("POST required.")
    form = ReceiptConditionForm(request.POST)
    if form.is_valid():
        try:
            services.confirm_receipt(
                user=request.user, relocation=reloc,
                condition_at_receipt=form.cleaned_data["condition_at_receipt"],
                receipt_notes=form.cleaned_data["receipt_notes"],
            )
            messages.success(request, "Receipt confirmed. Asset relocated.")
        except (PermissionDenied, ValidationError) as e:
            messages.error(request, str(e))
    return redirect("property:relocation_detail", pk=reloc.pk)


@login_required
@require_director
def relocation_resolve_exception(request, pk):
    reloc = get_object_or_404(AssetRelocation, pk=pk)
    if request.method != "POST":
        raise PermissionDenied("POST required.")
    form = RelocationExceptionResolveForm(request.POST)
    if form.is_valid():
        try:
            services.resolve_relocation_exception(
                user=request.user, relocation=reloc,
                action=form.cleaned_data["action"],
                cancellation_reason=form.cleaned_data.get("cancellation_reason", ""),
                new_recipient=form.cleaned_data.get("new_recipient"),
            )
            messages.success(request, "Exception resolved.")
        except (PermissionDenied, ValidationError) as e:
            messages.error(request, str(e))
    return redirect("property:approval_queue")


# ---------------------------------------------------------------------------
# Director reports + CSV export
# ---------------------------------------------------------------------------
@login_required
@require_director
def reports(request):
    form = ReportFilterForm(request.GET or None)
    qs = AssetRelocation.objects.select_related(
        "asset", "origin_area", "destination_area",
        "requested_by", "receiving_attendant", "approved_by", "received_by",
    ).order_by("-submitted_at")
    if form.is_valid():
        if form.cleaned_data.get("date_from"):
            qs = qs.filter(submitted_at__date__gte=form.cleaned_data["date_from"])
        if form.cleaned_data.get("date_to"):
            qs = qs.filter(submitted_at__date__lte=form.cleaned_data["date_to"])
        if form.cleaned_data.get("status"):
            qs = qs.filter(status=form.cleaned_data["status"])
        if form.cleaned_data.get("origin_area"):
            qs = qs.filter(origin_area=form.cleaned_data["origin_area"])
        if form.cleaned_data.get("destination_area"):
            qs = qs.filter(destination_area=form.cleaned_data["destination_area"])
    return render(request, "property/reports.html", {"form": form, "rows": qs[:500]})


@login_required
@require_director
def reports_csv(request):
    form = ReportFilterForm(request.GET or None)
    qs = AssetRelocation.objects.select_related(
        "asset", "origin_area", "destination_area",
        "requested_by", "receiving_attendant", "approved_by", "received_by",
    ).order_by("-submitted_at")
    if form.is_valid():
        if form.cleaned_data.get("date_from"):
            qs = qs.filter(submitted_at__date__gte=form.cleaned_data["date_from"])
        if form.cleaned_data.get("date_to"):
            qs = qs.filter(submitted_at__date__lte=form.cleaned_data["date_to"])
        if form.cleaned_data.get("status"):
            qs = qs.filter(status=form.cleaned_data["status"])

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="relocations_{timezone.now():%Y%m%d}.csv"'
    w = csv.writer(response)
    w.writerow([
        "Asset #", "Asset name", "Origin", "Destination",
        "Requested by", "Requested at", "Reason",
        "Condition before", "Director", "Decision", "Approved at",
        "Recipient", "Recipient decision", "Decline reason",
        "Condition at receipt", "Receipt notes", "Completed at",
        "Current location", "Status",
    ])
    for r in qs:
        w.writerow([
            r.asset.asset_number, r.asset.name,
            r.origin_area.code, r.destination_area.code,
            r.requested_by.username if r.requested_by_id else "",
            r.submitted_at, r.reason,
            r.condition_before,
            r.approved_by.username if r.approved_by_id else "",
            r.status, r.approved_at or "",
            r.receiving_attendant.username if r.receiving_attendant_id else "",
            r.recipient_decision, r.decline_reason,
            r.condition_at_receipt, r.receipt_notes,
            r.completed_at or "",
            r.asset.current_area.code,
            r.status,
        ])
    return response

# ===========================================================================
# STAFF AREA ACCESS (Director only)
# ===========================================================================
@login_required
@require_director
def staff_access(request):
    """Matrix of active staff x physical areas with checkboxes."""
    from accounts.models import Employee
    from .models import StaffAreaAccess
    staff = (Employee.objects.filter(
        user__is_active=True, user__is_active_staff=True)
        .select_related("user", "operating_point")
        .order_by("user__username"))
    areas = PhysicalArea.objects.filter(is_active=True).order_by("code")
    # Map of (employee_id, area_code) -> is_active
    existing = {}
    for sa in StaffAreaAccess.objects.select_related("area"):
        existing[(sa.employee_id, sa.area.code)] = sa.is_active

    matrix = []
    for emp in staff:
        row = []
        for area in areas:
            row.append({
                "area": area,
                "granted": existing.get((emp.pk, area.code), False),
            })
        matrix.append({"employee": emp, "cells": row})

    return render(request, "property/staff_access.html", {
        "areas": areas,
        "matrix": matrix,
    })


@login_required
@require_director
def staff_access_update(request):
    """Process form submission: per-(employee, area) grant/revoke."""
    if request.method != "POST":
        raise PermissionDenied("POST required.")
    from accounts.models import Employee
    from .models import StaffAreaAccess

    changed = 0
    for emp in Employee.objects.filter(
            user__is_active=True, user__is_active_staff=True):
        for area in PhysicalArea.objects.filter(is_active=True):
            checkbox_key = f"access_{emp.pk}_{area.code}"
            want = checkbox_key in request.POST
            existing = StaffAreaAccess.objects.filter(
                employee=emp, area=area).first()
            if want and existing is None:
                StaffAreaAccess.objects.create(
                    employee=emp, area=area,
                    granted_by=request.user, is_active=True)
                changed += 1
            elif not want and existing is not None and existing.is_active:
                existing.is_active = False
                existing.save(update_fields=["is_active"])
                changed += 1
            elif want and existing is not None and not existing.is_active:
                existing.is_active = True
                existing.save(update_fields=["is_active"])
                changed += 1

    messages.success(request, f"Updated {changed} grant(s).")
    return redirect("property:staff_access")

# ===========================================================================
# PDF + Excel exports (Director only)
# ===========================================================================
@login_required
@require_director
def reports_pdf(request):
    from io import BytesIO
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    form = ReportFilterForm(request.GET or None)
    qs = _filtered_relocations(form)

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    story = [Paragraph("Asset Relocations Report", styles["Title"]),
             Spacer(1, 12)]

    data = [["Asset #", "From", "To", "Requested", "Director",
             "Recipient", "Cond before", "Cond on receipt", "Completed", "Status"]]
    for r in qs[:1000]:
        data.append([
            r.asset.asset_number,
            r.origin_area.code,
            r.destination_area.code,
            r.submitted_at.strftime("%Y-%m-%d %H:%M"),
            r.approved_by.username if r.approved_by_id else "—",
            r.receiving_attendant.username if r.receiving_attendant_id else "—",
            r.condition_before,
            r.condition_at_receipt or "—",
            r.completed_at.strftime("%Y-%m-%d %H:%M") if r.completed_at else "—",
            r.status,
        ])
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
    ]))
    story.append(t)
    doc.build(story)

    pdf = buf.getvalue()
    buf.close()
    resp = HttpResponse(pdf, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="relocations_{timezone.now():%Y%m%d}.pdf"'
    return resp


@login_required
@require_director
def reports_xlsx(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    form = ReportFilterForm(request.GET or None)
    qs = _filtered_relocations(form)

    wb = Workbook()
    ws = wb.active
    ws.title = "Relocations"
    headers = ["Asset #", "Asset name", "From", "To",
               "Requested by", "Requested at", "Reason",
               "Condition before", "Director", "Decision",
               "Approved at", "Recipient", "Recipient decision",
               "Decline reason", "Condition at receipt",
               "Receipt notes", "Completed at", "Current location", "Status"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for r in qs[:5000]:
        ws.append([
            r.asset.asset_number, r.asset.name,
            r.origin_area.code, r.destination_area.code,
            r.requested_by.username if r.requested_by_id else "",
            r.submitted_at.replace(tzinfo=None) if r.submitted_at else "",
            r.reason,
            r.condition_before,
            r.approved_by.username if r.approved_by_id else "",
            r.status,
            r.approved_at.replace(tzinfo=None) if r.approved_at else "",
            r.receiving_attendant.username if r.receiving_attendant_id else "",
            r.recipient_decision or "",
            r.decline_reason or "",
            r.condition_at_receipt or "",
            r.receipt_notes or "",
            r.completed_at.replace(tzinfo=None) if r.completed_at else "",
            r.asset.current_area.code,
            r.status,
        ])
    buf = BytesIO()
    wb.save(buf)
    xlsx = buf.getvalue()
    buf.close()
    resp = HttpResponse(
        xlsx,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="relocations_{timezone.now():%Y%m%d}.xlsx"'
    return resp


def _filtered_relocations(form):
    qs = AssetRelocation.objects.select_related(
        "asset", "origin_area", "destination_area",
        "requested_by", "receiving_attendant", "approved_by", "received_by",
    ).order_by("-submitted_at")
    if form.is_valid():
        if form.cleaned_data.get("date_from"):
            qs = qs.filter(submitted_at__date__gte=form.cleaned_data["date_from"])
        if form.cleaned_data.get("date_to"):
            qs = qs.filter(submitted_at__date__lte=form.cleaned_data["date_to"])
        if form.cleaned_data.get("status"):
            qs = qs.filter(status=form.cleaned_data["status"])
        if form.cleaned_data.get("origin_area"):
            qs = qs.filter(origin_area=form.cleaned_data["origin_area"])
        if form.cleaned_data.get("destination_area"):
            qs = qs.filter(destination_area=form.cleaned_data["destination_area"])
    return qs

# ===========================================================================
# Popup notification API
# ===========================================================================
from django.http import JsonResponse
from django.views.decorators.http import require_POST
from django.core.cache import cache


@login_required
@require_active_staff
def api_pending_property(request):
    """Return user-specific property action items the user hasn't dismissed."""
    user = request.user
    dismissed = set(cache.get(f"property_dismissed_{user.pk}", []))
    out = []

    if is_director(user):
        for r in (AssetRelocation.objects
                  .filter(status=AssetRelocation.Status.PENDING_APPROVAL)
                  .select_related("asset", "origin_area", "destination_area",
                                  "requested_by", "receiving_attendant")
                  .order_by("submitted_at")[:20]):
            key = f"reloc-pending-{r.pk}"
            if key in dismissed:
                continue
            out.append({
                "id": key,
                "title": "Approval required",
                "asset": f"{r.asset.asset_number} — {r.asset.name}",
                "from_area": r.origin_area.code,
                "to_area": r.destination_area.code,
                "message": f"Requested by {r.requested_by.username}",
                "created_at": r.submitted_at.isoformat(),
                "url": f"/property/relocations/{r.pk}/",
                "ack_url": f"/property/api/ack/{key}/",
            })
    else:
        for r in (AssetRelocation.objects
                  .filter(receiving_attendant=user,
                          status__in=[AssetRelocation.Status.APPROVED,
                                      AssetRelocation.Status.AWAITING_RECIPIENT])
                  .select_related("asset", "origin_area", "destination_area")
                  .order_by("submitted_at")[:20]):
            key = f"reloc-recipient-{r.pk}"
            if key in dismissed:
                continue
            out.append({
                "id": key,
                "title": "Transfer awaiting your receipt",
                "asset": f"{r.asset.asset_number} — {r.asset.name}",
                "from_area": r.origin_area.code,
                "to_area": r.destination_area.code,
                "message": "Approve, then confirm receipt",
                "created_at": r.submitted_at.isoformat(),
                "url": f"/property/relocations/{r.pk}/",
                "ack_url": f"/property/api/ack/{key}/",
            })

    return JsonResponse({"alerts": out})


@login_required
@require_POST
def api_ack_property(request, key):
    lst = cache.get(f"property_dismissed_{request.user.pk}", [])
    if key not in lst:
        lst.append(key)
    cache.set(f"property_dismissed_{request.user.pk}", lst, timeout=60 * 60 * 24 * 7)
    return JsonResponse({"ok": True})


# ---------------------------------------------------------------------------
# Movement report
# ---------------------------------------------------------------------------
from datetime import timedelta as _pm_td


@login_required
@require_director
def movement_report(request):
    try:
        days = int(request.GET.get("days", 30) or 30)
    except (TypeError, ValueError):
        days = 30
    days = max(1, min(days, 365))

    status_filter = (request.GET.get("status") or "").strip().upper()

    since = timezone.now() - _pm_td(days=days)

    qs = (AssetRelocation.objects
          .filter(submitted_at__gte=since)
          .select_related(
              "asset", "origin_area", "destination_area",
              "requested_by", "receiving_attendant",
          )
          .order_by("-submitted_at"))

    if status_filter:
        qs = qs.filter(status=status_filter)

    entries = list(qs[:500])

    # --- summary counts by status
    summary = {}
    for e in entries:
        summary[e.status] = summary.get(e.status, 0) + 1

    return render(request, "property/movement_report.html", {
        "entries": entries,
        "entry_count": len(entries),
        "days": days,
        "status_filter": status_filter,
        "status_choices": AssetRelocation.Status.choices,
        "summary": summary,
    })
